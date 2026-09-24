#!/usr/bin/env python3
"""Finalize hashes and validate the generated Wathiq sample corpus."""

from __future__ import annotations

import csv
import email
from email import policy
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
from xml.etree import ElementTree as ET


def count_words(text: str) -> int:
    return len(text.split())


def extract_body(root: Path, row: dict[str, str]) -> str | None:
    path = root / row["path"]
    extension = row["format"]
    if extension == "txt":
        return " ".join(path.read_text(encoding="utf-8").splitlines()[4:])
    if extension == "html":
        from html.parser import HTMLParser
        class Parser(HTMLParser):
            def __init__(self):
                super().__init__(); self.parts = []; self.in_p = False
            def handle_starttag(self, tag, attrs): self.in_p = tag == "p"
            def handle_endtag(self, tag):
                if tag == "p": self.in_p = False
            def handle_data(self, data):
                if self.in_p: self.parts.append(data)
        parser = Parser(); parser.feed(path.read_text(encoding="utf-8")); return " ".join(parser.parts)
    if extension == "docx":
        doc = Document(path)
        return " ".join(p.text for p in doc.paragraphs[3:])
    if extension == "xlsx":
        wb = load_workbook(path, read_only=True, data_only=False)
        ws = wb["Document"]
        text = " ".join(str(ws.cell(r, 1).value or "") for r in range(5, 15))
        wb.close(); return text
    if extension == "pdf":
        # Text extraction is checked for presence and bilingual markers; exact
        # word count is authoritative in the source manifest because PDF text
        # extraction may segment punctuation and Arabic differently.
        return " ".join((page.extract_text() or "") for page in PdfReader(path).pages)
    if extension == "eml":
        message = email.message_from_bytes(path.read_bytes(), policy=policy.default)
        part = message.get_body(preferencelist=("plain",))
        return part.get_content() if part is not None else ""
    if extension == "xml":
        return ET.parse(path).getroot().findtext("body") or ""
    if extension == "md":
        lines = path.read_text(encoding="utf-8").splitlines()
        return " ".join(line for line in lines[6:] if line.strip())
    if extension == "pptx":
        presentation = Presentation(path)
        parts = []
        for slide in presentation.slides:
            candidates = [shape.text for shape in slide.shapes if hasattr(shape, "text_frame")]
            body = max(candidates, key=count_words, default="")
            parts.append(body)
        return " ".join(parts)
    return None


def main() -> None:
    root = Path(sys.argv[1]).resolve()
    manifest = root / "manifest.csv"
    with manifest.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    errors = []
    counts = Counter(row["format"] for row in rows)
    expected = {
        "docx": 143, "xlsx": 143, "txt": 143, "pdf": 143, "png": 143,
        "jpeg": 143, "html": 142, "eml": 40, "msg": 40, "xml": 40,
        "pptx": 40, "md": 40,
    }
    if len(rows) != 1200 or dict(counts) != expected:
        errors.append(f"unexpected distribution: total={len(rows)} formats={dict(counts)}")

    for row in rows:
        path = root / row["path"]
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"missing or empty: {row['path']}")
            continue
        data = path.read_bytes()
        row["size_bytes"] = str(len(data))
        row["sha256"] = hashlib.sha256(data).hexdigest()
        body = extract_body(root, row)
        if row["format"] == "msg":
            if not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
                errors.append(f"invalid MSG compound-file signature: {row['path']}")
            for term_field in ("rare_english", "rare_arabic", "unique_marker"):
                encoded = row[term_field].encode("utf-16le")
                if encoded not in data:
                    errors.append(f"missing MSG {term_field}: {row['path']}")
        if body is not None:
            if row["format"] in {"docx", "xlsx", "txt", "html", "eml", "xml", "md", "pptx"} and count_words(body) != 200:
                errors.append(f"word count {count_words(body)}: {row['path']}")
            for term_field in ("rare_english", "rare_arabic", "unique_marker"):
                if row[term_field] not in body:
                    errors.append(f"missing {term_field}: {row['path']}")

    with manifest.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=rows[0].keys())
        writer.writeheader(); writer.writerows(rows)
    summary = {
        "total": len(rows), "formats": dict(counts),
        "language_modes": dict(Counter(row["language_mode"] for row in rows)),
        "topics": dict(Counter(row["topic"] for row in rows)),
        "errors": errors,
    }
    (root / "validation-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if errors:
        raise SystemExit("\n".join(errors[:30]))
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
