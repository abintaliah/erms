#!/usr/bin/env python3
"""Generate Wathiq's deterministic bilingual full-text-search sample corpus."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


FORMATS = (
    ("docx", 143),
    ("xlsx", 143),
    ("txt", 143),
    ("pdf", 143),
    ("png", 143),
    ("jpeg", 143),
    ("html", 142),
)

TOPICS = (
    "finance",
    "human-resources",
    "general-administration",
    "asset-management",
    "records-management",
    "project-management",
    "information-technology",
    "procurement",
)

TOPIC_LABELS = {
    "finance": ("Quarterly Finance Review", "مراجعة مالية ربع سنوية"),
    "human-resources": ("Workforce Planning Note", "مذكرة تخطيط القوى العاملة"),
    "general-administration": ("Administrative Services Update", "تحديث الخدمات الإدارية"),
    "asset-management": ("Asset Register Review", "مراجعة سجل الأصول"),
    "records-management": ("Records Classification Review", "مراجعة تصنيف السجلات"),
    "project-management": ("Project Delivery Update", "تحديث تنفيذ المشروع"),
    "information-technology": ("Technology Operations Review", "مراجعة عمليات التقنية"),
    "procurement": ("Procurement Planning Note", "مذكرة تخطيط المشتريات"),
}

ENGLISH_RARE = (
    "palimpsest", "quincunx", "zeugma", "syzygy", "susurrus", "velleity",
    "anfractuous", "crepuscular", "psithurism", "ultracrepidarian",
)
ARABIC_RARE = (
    "الاستيثاق", "التغاير", "الاستدلال", "الاستعصاء", "التماثل",
    "الاستقراء", "التبويب", "المضاهاة", "الاستدراك", "الاستدامة",
)

ENGLISH_SENTENCES = {
    "finance": [
        "The finance team reviewed monthly expenditure, revenue timing, and cash forecasts before preparing the next management report.",
        "Analysts reconciled supplier invoices with approved commitments and documented each variance that required follow-up by the budget owner.",
        "The revised forecast preserves contingency funding while directing available resources toward services with confirmed operational demand.",
        "Managers will compare actual spending with the approved budget and explain material changes during the quarterly review meeting.",
    ],
    "human-resources": [
        "Human resources reviewed recruitment demand, employee development plans, and leave coverage for the coming operational period.",
        "Managers confirmed that position descriptions reflect current responsibilities and that interview records follow the approved retention schedule.",
        "The workforce plan balances essential vacancies with training commitments and provides clear ownership for every pending action.",
        "Employee records remain restricted to authorized staff, with access reviews scheduled before the next reporting cycle.",
    ],
    "general-administration": [
        "The administration team coordinated meeting rooms, correspondence, travel requests, and shared services for the coming month.",
        "Staff recorded each request with an owner, due date, and supporting reference so outstanding work remains visible.",
        "The updated procedure reduces duplicate approvals and keeps evidence beside the transaction that produced it.",
        "Supervisors will review delayed requests weekly and resolve dependencies with the responsible business unit.",
    ],
    "asset-management": [
        "The asset team inspected assigned equipment and compared physical labels with the central register and custody records.",
        "Each discrepancy includes a location, responsible custodian, review date, and documented action for correction.",
        "Planned maintenance gives priority to safety, service continuity, and equipment approaching the end of its useful life.",
        "Disposed assets require approved evidence before their records can be closed in the management system.",
    ],
    "records-management": [
        "Records staff reviewed file classifications, retention periods, access controls, and transfer readiness across active business folders.",
        "Every digital component must remain linked to its authoritative record and retain a verifiable content checksum.",
        "The review identified duplicate working copies and assigned owners to confirm disposition without removing evidence under legal hold.",
        "Metadata corrections will follow the approved audit process and preserve the history of each governed change.",
    ],
    "project-management": [
        "The project team reviewed milestones, dependencies, risks, and decisions before updating the delivery schedule.",
        "Workstream leads confirmed completed tasks with evidence and identified constraints that could affect the next reporting period.",
        "The current plan protects critical testing time while allowing lower priority activities to move within agreed tolerances.",
        "Sponsors will receive a concise status report showing progress, unresolved decisions, and accountable action owners.",
    ],
    "information-technology": [
        "The technology team reviewed service incidents, backup results, access changes, and planned maintenance for core systems.",
        "Engineers documented the cause of recurring alerts and assigned corrective actions with measurable completion criteria.",
        "Security updates will be tested in the controlled environment before deployment to production services.",
        "Operational dashboards distinguish customer impact from internal warnings and retain evidence for later analysis.",
    ],
    "procurement": [
        "The procurement team reviewed purchase requests, evaluation records, contract dates, and supplier performance information.",
        "Each sourcing decision must retain the approved specification, evaluation evidence, and documented authorization trail.",
        "Contract owners will monitor delivery dates and resolve discrepancies before accepting invoices for payment.",
        "The updated plan groups related demand where consolidation improves value without weakening operational accountability.",
    ],
}

ARABIC_SENTENCES = {
    "finance": [
        "راجع فريق المالية المصروفات الشهرية وتوقيت الإيرادات وتوقعات التدفق النقدي قبل إعداد التقرير الإداري القادم.",
        "طابق المحللون فواتير الموردين مع الالتزامات المعتمدة ووثقوا كل فرق يحتاج إلى متابعة من مسؤول الميزانية.",
        "يحافظ التوقع المحدث على مخصص الطوارئ ويوجه الموارد المتاحة إلى الخدمات ذات الحاجة التشغيلية المؤكدة.",
        "سيقارن المديرون الإنفاق الفعلي بالميزانية المعتمدة ويشرحون التغيرات المهمة خلال اجتماع المراجعة الربع سنوية.",
    ],
    "human-resources": [
        "راجعت الموارد البشرية احتياجات التوظيف وخطط تطوير الموظفين وتغطية الإجازات للفترة التشغيلية القادمة.",
        "أكد المديرون أن الأوصاف الوظيفية تعكس المسؤوليات الحالية وأن سجلات المقابلات تتبع مدة الحفظ المعتمدة.",
        "توازن خطة القوى العاملة بين الشواغر الأساسية والتزامات التدريب وتحدد مسؤولية واضحة لكل إجراء معلق.",
        "تبقى سجلات الموظفين مقيدة بالمخولين مع جدولة مراجعة الصلاحيات قبل دورة التقارير القادمة.",
    ],
    "general-administration": [
        "نسق فريق الإدارة قاعات الاجتماعات والمراسلات وطلبات السفر والخدمات المشتركة للشهر القادم.",
        "سجل الموظفون كل طلب مع المسؤول وتاريخ الاستحقاق والمرجع الداعم حتى يبقى العمل المعلق واضحا.",
        "يقلل الإجراء المحدث الموافقات المكررة ويحفظ الأدلة بجانب المعاملة التي أنتجتها.",
        "سيراجع المشرفون الطلبات المتأخرة أسبوعيا ويعالجون الاعتماديات مع وحدة الأعمال المسؤولة.",
    ],
    "asset-management": [
        "فحص فريق الأصول المعدات المسندة وقارن الملصقات الفعلية بالسجل المركزي ووثائق العهدة.",
        "يتضمن كل اختلاف الموقع وأمين العهدة وتاريخ المراجعة والإجراء الموثق المطلوب للتصحيح.",
        "تعطي الصيانة المخططة الأولوية للسلامة واستمرار الخدمة والمعدات التي تقترب من نهاية عمرها النافع.",
        "تتطلب الأصول المستبعدة دليلا معتمدا قبل إغلاق سجلاتها في نظام الإدارة.",
    ],
    "records-management": [
        "راجع موظفو السجلات التصنيفات ومدد الحفظ وضوابط الوصول وجاهزية التحويل في ملفات الأعمال النشطة.",
        "يجب أن يبقى كل مكون رقمي مرتبطا بسجله الرسمي وأن يحتفظ ببصمة محتوى قابلة للتحقق.",
        "حددت المراجعة نسخ العمل المكررة وكلفت المسؤولين بتأكيد التصرف دون إزالة دليل خاضع للحجز القانوني.",
        "ستتبع تصحيحات البيانات الوصفية مسار التدقيق المعتمد وتحافظ على تاريخ كل تغيير محكوم.",
    ],
    "project-management": [
        "راجع فريق المشروع المراحل والاعتماديات والمخاطر والقرارات قبل تحديث جدول التنفيذ.",
        "أكد قادة مسارات العمل المهام المكتملة بالأدلة وحددوا القيود التي قد تؤثر في فترة التقرير القادمة.",
        "تحمي الخطة الحالية وقت الاختبار الحرج وتسمح بتحريك الأنشطة الأقل أولوية ضمن الحدود المتفق عليها.",
        "سيتلقى الرعاة تقرير حالة موجزا يوضح التقدم والقرارات غير المحسومة والمسؤولين عن الإجراءات.",
    ],
    "information-technology": [
        "راجع فريق التقنية حوادث الخدمة ونتائج النسخ الاحتياطي وتغييرات الوصول والصيانة المخططة للأنظمة الأساسية.",
        "وثق المهندسون سبب التنبيهات المتكررة وأسندوا إجراءات تصحيحية بمعايير إنجاز قابلة للقياس.",
        "ستختبر التحديثات الأمنية في البيئة المنضبطة قبل نشرها على خدمات الإنتاج.",
        "تميز لوحات التشغيل أثر المستخدم عن التحذيرات الداخلية وتحفظ الأدلة للتحليل اللاحق.",
    ],
    "procurement": [
        "راجع فريق المشتريات طلبات الشراء وسجلات التقييم وتواريخ العقود ومعلومات أداء الموردين.",
        "يجب أن يحتفظ كل قرار توريد بالمواصفات المعتمدة وأدلة التقييم ومسار التفويض الموثق.",
        "سيراقب مسؤولو العقود تواريخ التسليم ويعالجون الاختلافات قبل قبول الفواتير للدفع.",
        "تجمع الخطة المحدثة الطلبات المرتبطة عندما يحسن الدمج القيمة دون إضعاف المساءلة التشغيلية.",
    ],
}


def words(text: str) -> list[str]:
    return text.split()


def make_body(index: int, topic: str, mode: str) -> tuple[str, str, str]:
    rare_en = ENGLISH_RARE[index % len(ENGLISH_RARE)]
    rare_ar = ARABIC_RARE[(index * 3) % len(ARABIC_RARE)]
    marker = f"wathiqrare{index:04d}"
    en = ENGLISH_SENTENCES[topic]
    ar = ARABIC_SENTENCES[topic]
    if mode == "english":
        sequence = [en[0], en[1], en[2], en[3], ar[0]]
    elif mode == "arabic":
        sequence = [ar[0], ar[1], ar[2], ar[3], en[0]]
    else:
        sequence = [en[0], ar[0], en[1], ar[1], en[2], ar[2]]

    prefix = (
        f"Reference {marker} uses the uncommon term {rare_en}. "
        f"وتتضمن المذكرة المصطلح النادر {rare_ar}. "
    )
    material = prefix
    cycle = 0
    while len(words(material)) < 200:
        sentence = sequence[(index + cycle) % len(sequence)]
        material += sentence + " "
        cycle += 1
    tokens = words(material)[:200]
    # The source sentences are repeated as complete units; replace a truncated
    # tail with a grammatical bilingual closing of the same token count.
    closing = words("The responsible team will retain evidence and review progress. وسيحفظ الفريق المسؤول الأدلة ويراجع التقدم.")
    if len(tokens) >= len(closing):
        tokens[-len(closing):] = closing
    assert len(tokens) == 200
    return " ".join(tokens), rare_en, rare_ar


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")


def write_docx(path: Path, title: str, arabic_title: str, reference: str, body: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run(title), 16, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(p.add_run(arabic_title), 14, True)
    p = doc.add_paragraph()
    set_run_font(p.add_run(f"Synthetic reference {reference}"), 9)

    body_tokens = words(body)
    for start in range(0, len(body_tokens), 50):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.04
        set_run_font(p.add_run(" ".join(body_tokens[start:start + 50])), 10.5)

    doc.core_properties.title = title
    doc.core_properties.subject = "Synthetic bilingual Wathiq full-text-search fixture"
    doc.core_properties.author = "Wathiq Sample Corpus Generator"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def write_txt(path: Path, title: str, arabic_title: str, reference: str, body: str) -> None:
    path.write_text(f"{title}\n{arabic_title}\nSynthetic reference {reference}\n\n{body}\n", encoding="utf-8")


def write_html(path: Path, title: str, arabic_title: str, reference: str, body: str) -> None:
    import html
    paragraphs = [" ".join(words(body)[i:i + 50]) for i in range(0, 200, 50)]
    rendered = "\n".join(f"<p>{html.escape(p)}</p>" for p in paragraphs)
    path.write_text(
        "<!doctype html>\n<html lang=\"en\"><head><meta charset=\"utf-8\">"
        f"<title>{html.escape(title)}</title><style>body{{font-family:Arial,sans-serif;max-width:760px;margin:48px auto;line-height:1.55;color:#172033}}"
        "h1{font-size:24px;margin-bottom:4px}h2{font-size:20px;text-align:right;margin-top:0}"
        ".ref{color:#526174;font-size:13px}p{font-size:15px}</style></head><body>"
        f"<h1>{html.escape(title)}</h1><h2 dir=\"rtl\">{html.escape(arabic_title)}</h2>"
        f"<div class=\"ref\">Synthetic reference {html.escape(reference)}</div>{rendered}</body></html>\n",
        encoding="utf-8",
    )


def convert_docx_batch(source_dir: Path, pdf_dir: Path, soffice: str) -> None:
    files = sorted(source_dir.glob("*.docx"))
    pdf_dir.mkdir(parents=True, exist_ok=True)
    for batch_no, start in enumerate(range(0, len(files), 40)):
        profile = source_dir.parent / f"lo-profile-{batch_no}"
        command = [
            soffice,
            "--headless",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to", "pdf",
            "--outdir", str(pdf_dir),
            *[str(p) for p in files[start:start + 40]],
        ]
        result = subprocess.run(command, capture_output=True, text=True)
        if result.returncode:
            raise RuntimeError(f"LibreOffice conversion failed: {result.stderr}")
    missing = [p.stem for p in files if not (pdf_dir / f"{p.stem}.pdf").exists()]
    if missing:
        raise RuntimeError(f"LibreOffice did not produce {len(missing)} PDFs: {missing[:5]}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--soffice", required=True)
    parser.add_argument("--pdftoppm", required=True)
    args = parser.parse_args()
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=True)
    for extension, _ in FORMATS:
        (root / extension).mkdir(exist_ok=True)

    specs = []
    number = 1
    for extension, count in FORMATS:
        for local_index in range(count):
            topic = TOPICS[(number - 1) % len(TOPICS)]
            mode = ("english", "arabic", "mixed")[(local_index + FORMATS.index((extension, count))) % 3]
            title, arabic_title = TOPIC_LABELS[topic]
            body, rare_en, rare_ar = make_body(number, topic, mode)
            stem = f"sample-{number:04d}-{topic}-{mode}"
            specs.append({
                "id": number,
                "format": extension,
                "path": f"{extension}/{stem}.{extension}",
                "stem": stem,
                "topic": topic,
                "language_mode": mode,
                "title": title,
                "arabic_title": arabic_title,
                "reference": f"WTHQ-{number:04d}",
                "body": body,
                "body_word_count": len(words(body)),
                "rare_english": rare_en,
                "rare_arabic": rare_ar,
                "unique_marker": f"wathiqrare{number:04d}",
            })
            number += 1
    assert len(specs) == 1000
    assert all(s["body_word_count"] == 200 for s in specs)

    with tempfile.TemporaryDirectory(prefix="wathiq-corpus-") as temporary:
        temp = Path(temporary)
        source_docs = temp / "source-docx"
        rendered_pdfs = temp / "rendered-pdf"
        source_docs.mkdir()

        for spec in specs:
            destination = root / spec["path"]
            ext = spec["format"]
            if ext == "txt":
                write_txt(destination, spec["title"], spec["arabic_title"], spec["reference"], spec["body"])
            elif ext == "html":
                write_html(destination, spec["title"], spec["arabic_title"], spec["reference"], spec["body"])
            elif ext in {"docx", "pdf", "png", "jpeg"}:
                docx_path = destination if ext == "docx" else source_docs / f"{spec['stem']}.docx"
                write_docx(docx_path, spec["title"], spec["arabic_title"], spec["reference"], spec["body"])

        # Render all final DOCX files too, providing a common visual-QA source.
        for spec in specs:
            if spec["format"] == "docx":
                shutil.copy2(root / spec["path"], source_docs / f"{spec['stem']}.docx")
        convert_docx_batch(source_docs, rendered_pdfs, args.soffice)

        for spec in specs:
            ext = spec["format"]
            source_pdf = rendered_pdfs / f"{spec['stem']}.pdf"
            destination = root / spec["path"]
            if ext == "pdf":
                shutil.copy2(source_pdf, destination)
            elif ext in {"png", "jpeg"}:
                output_prefix = destination.with_suffix("")
                command = [args.pdftoppm, "-f", "1", "-singlefile", "-r", "130", f"-{ext}", str(source_pdf), str(output_prefix)]
                result = subprocess.run(command, capture_output=True, text=True)
                if result.returncode:
                    raise RuntimeError(f"Raster conversion failed for {destination}: {result.stderr}")
                if ext == "jpeg":
                    poppler_path = output_prefix.with_suffix(".jpg")
                    if poppler_path.exists():
                        poppler_path.replace(destination)

        qa_dir = root / ".qa-rendered-pdf"
        qa_dir.mkdir(exist_ok=True)
        for spec in specs:
            if spec["format"] == "docx":
                shutil.copy2(rendered_pdfs / f"{spec['stem']}.pdf", qa_dir / f"{spec['stem']}.pdf")

    xlsx_specs = [s for s in specs if s["format"] == "xlsx"]
    (root / ".xlsx-payload.json").write_text(json.dumps(xlsx_specs, ensure_ascii=False), encoding="utf-8")

    rows = []
    for spec in specs:
        path = root / spec["path"]
        if spec["format"] == "xlsx":
            sha256 = "PENDING_XLSX_GENERATION"
            size = 0
        else:
            data = path.read_bytes()
            sha256 = hashlib.sha256(data).hexdigest()
            size = len(data)
        rows.append({
            "id": spec["id"], "format": spec["format"], "path": spec["path"],
            "topic": spec["topic"], "language_mode": spec["language_mode"],
            "body_word_count": spec["body_word_count"], "rare_english": spec["rare_english"],
            "rare_arabic": spec["rare_arabic"], "unique_marker": spec["unique_marker"],
            "size_bytes": size, "sha256": sha256,
        })
    with (root / "manifest.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)
    print(json.dumps({"files_prepared": len(specs), "formats": dict(Counter(s["format"] for s in specs))}))


if __name__ == "__main__":
    main()
