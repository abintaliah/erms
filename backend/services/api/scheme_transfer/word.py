"""On-demand Word renderer; document dependencies stay off page-load paths."""

import io
import unicodedata
from datetime import datetime, timezone

from babel import Locale, UnknownLocaleError
from babel.dates import format_datetime

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_DIRECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..entity_localization import localized_projection
from .codec import TransferError, ordered
from .messages import TRANSFER_MESSAGE_KEYS


def render_word(connection, package, language):
    config = connection.execute(
        "SELECT * FROM supported_languages WHERE language_tag=%s AND is_enabled",
        (language,),
    ).fetchone()
    if not config:
        raise TransferError("unsupported_language")
    # No global catalogue cache: query published wording inside this export's
    # snapshot so a freshly published translation is immediately reflected.
    rows = connection.execute(
        """SELECT d.message_key,COALESCE(t.published_text,b.published_text,d.default_text) AS text
      FROM ui_message_definitions d
      LEFT JOIN ui_message_translations t ON t.message_key=d.message_key AND t.language_tag=%s AND NOT t.needs_review
      LEFT JOIN ui_message_translations b ON b.message_key=d.message_key AND b.language_tag=%s AND NOT b.needs_review
      WHERE NOT d.is_deprecated""",
        (language, language.split("-")[0]),
    ).fetchall()
    messages = {r["message_key"]: r["text"] for r in rows}

    def label(key):
        return messages[TRANSFER_MESSAGE_KEYS[key]]

    rtl = config["direction"] == "rtl"
    locale_name = (config.get("formatting_config") or {}).get("locale", language)
    if not isinstance(locale_name, str):
        locale_name = language
    try:
        locale = Locale.parse(locale_name.replace("-", "_"))
    except (ValueError, UnknownLocaleError):
        try:
            locale = Locale.parse(language.split("-")[0])
        except (ValueError, UnknownLocaleError):
            locale = Locale.parse("en")

    def value(item, key=None):
        if item is None:
            return label("none")
        if isinstance(item, bool):
            return label("yes" if item else "no")
        if key and (key.startswith("date_") or key == "exported_at"):
            timestamp = datetime.fromisoformat(str(item).replace("Z", "+00:00"))
            return format_datetime(
                timestamp, "d MMMM y, HH:mm 'UTC'", tzinfo=timezone.utc, locale=locale
            )
        return str(item)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(13.5), Inches(8.5)
    section.left_margin = section.right_margin = Inches(0.55)
    section.top_margin = section.bottom_margin = Inches(0.5)
    style = document.styles["Normal"]
    style.font.name = "Changa"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(4)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), language)
    lang.set(qn("w:bidi"), language)
    style.element.get_or_add_rPr().append(lang)
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        selected_style = document.styles[name]
        selected_style.font.name = "Changa"
        selected_style.font.size = Pt(10)
        properties = selected_style.element.get_or_add_rPr()
        fonts = properties.find(qn("w:rFonts"))
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attribute), "Changa")
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attribute), None)
        complex_size = properties.find(qn("w:szCs"))
        if complex_size is None:
            complex_size = OxmlElement("w:szCs")
            properties.append(complex_size)
        complex_size.set(qn("w:val"), "20")

    def write_text(paragraph, text):
        # Use native Word runs, never inject Unicode direction controls into
        # document text. Strong letters/numbers select the run direction;
        # punctuation follows its surrounding run. New lines start afresh.
        for line_index, line in enumerate(str(text).split("\n")):
            if line_index:
                paragraph.add_run().add_break()
            run_direction = rtl
            buffer = ""
            for character in line:
                bidi_class = unicodedata.bidirectional(character)
                next_direction = (
                    True
                    if bidi_class in ("R", "AL")
                    else False if bidi_class in ("L", "EN", "AN") else run_direction
                )
                if buffer and next_direction != run_direction:
                    paragraph.add_run(buffer).font.rtl = run_direction
                    buffer = ""
                buffer += character
                run_direction = next_direction
            if buffer:
                paragraph.add_run(buffer).font.rtl = run_direction

    def direction(paragraph, depth=0):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph._p.get_or_add_pPr().find(qn("w:jc")).set(qn("w:val"), "start")
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1" if rtl else "0")
        paragraph._p.get_or_add_pPr().append(bidi)
        # Cap physical indentation only; explicit level and parent remain visible.
        # Logical start follows bidi direction. A physical w:right indent is
        # interpreted as the end edge by some RTL renderers (including LO).
        indent = OxmlElement("w:ind")
        indent.set(qn("w:start"), str(Inches(min(depth * 0.12, 0.8)).twips))
        indent.set(qn("w:end"), "0")
        paragraph._p.get_or_add_pPr().append(indent)

    def paragraph(text, style=None):
        p = document.add_paragraph(style=style)
        direction(p)
        write_text(p, text)
        return p

    def cell(cell, text, depth=0):
        cell.text = ""
        p = cell.paragraphs[0]
        direction(p, depth)
        write_text(p, text)

    def table(headers, widths):
        result = document.add_table(rows=1, cols=len(headers))
        result.style = "Table Grid"
        result.autofit = False
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement("w:" + edge)
            for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                element.set(qn("w:" + key), val)
            borders.append(element)
        result._tbl.tblPr.append(borders)
        result.table_direction = (
            WD_TABLE_DIRECTION.RTL if rtl else WD_TABLE_DIRECTION.LTR
        )
        for column, width in zip(result.columns, widths):
            column.width = Inches(width)
        for target, width in zip(result.rows[0].cells, widths):
            target.width = Inches(width)
        repeat = OxmlElement("w:tblHeader")
        result.rows[0]._tr.get_or_add_trPr().append(repeat)
        for target, text in zip(result.rows[0].cells, headers):
            cell(target, text)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DBEAFE")
            target._tc.get_or_add_tcPr().append(shade)
        return result

    disposition_keys = {
        "destruction": "webui.field_input.select.destruction_2711246f",
        "transfer_to_external_archive": "webui.field_input.select.permanent_preservation_external_archive_53d3f9f7",
        "selective_preservation": "webui.field_input.select.selective_preservation_04d0086f",
        "retain_as_local_archives": "webui.field_input.select.retain_as_local_archives_8d22b8ed",
    }

    def disposition(item):
        return messages.get(disposition_keys.get(item), value(item))

    def metadata(entity, omit=()):
        result = []
        for key, item in entity.items():
            if key in omit or key in ("classifications", "retention_rule"):
                continue
            if key == "translations":
                for tag, translated in (item or {}).items():
                    for field, text in translated.items():
                        result.append(f"{tag} · {label(field)}: {text}")
            else:
                if key == "final_disposition":
                    item = disposition(item)
                result.append(f"{label(key)}: {value(item, key)}")
        return "\n".join(result)

    scheme = package["data"]["scheme"]
    manifest = package["manifest"]
    title = localized_projection(scheme, language, "title")["title"]
    paragraph(title, "Title")
    paragraph(
        f"{label('exported_at')}: {value(manifest['exported_at'], 'exported_at')}\n{label('exported_by')}: {value(manifest['exported_by']['display_name'] or manifest['exported_by']['username'])}\n{label('database_name')}: {value(manifest['source']['database_name'])}"
    )
    summary = table([label("metadata"), label("value")], [2.2, 10.2])
    for key, item in scheme.items():
        if key == "classifications":
            continue
        cells = summary.add_row().cells
        cell(cells[0], label(key))
        cell(
            cells[1],
            (
                metadata({"translations": item})
                if key == "translations"
                else value(item, key)
            ),
        )
    paragraph(label("legend"))
    depths, effective = {}, {}
    current_table = None
    for item in ordered(scheme["classifications"]):
        parent = item["parent_code"]
        depth = depths.get(parent, -1) + 1
        depths[item["code"]] = depth
        if parent is None:
            paragraph(
                f"{item['code']} · {localized_projection(item, language, 'title')['title']}",
                "Heading 1",
            )
            current_table = table(
                [
                    label(k)
                    for k in (
                        "code",
                        "level",
                        "title",
                        "current_period_years",
                        "intermediate_period_years",
                        "final_disposition",
                        "metadata",
                        "retention_rule",
                    )
                ],
                [0.85, 0.8, 2.1, 0.8, 0.85, 1.65, 3.15, 2.2],
            )
        rule = item["retention_rule"]
        effective[item["code"]] = (
            (item["code"], rule) if rule else effective.get(parent)
        )
        governing = effective[item["code"]]
        retention = (
            (label("explicit") if rule else label("inherited") + " " + governing[0])
            + "\n"
            + metadata(
                governing[1],
                (
                    "current_period_years",
                    "intermediate_period_years",
                    "final_disposition",
                ),
            )
            if governing
            else label("none")
        )
        cells = current_table.add_row().cells
        values = [
            item["code"],
            str(depth),
            localized_projection(item, language, "title")["title"],
            value(governing[1]["current_period_years"] if governing else None),
            value(governing[1]["intermediate_period_years"] if governing else None),
            disposition(governing[1]["final_disposition"] if governing else None),
            metadata(item, ("code",)),
            retention,
        ]
        for index, (target, text) in enumerate(zip(cells, values)):
            cell(target, text, depth if index == 2 else 0)
            shade = OxmlElement("w:shd")
            shade.set(
                qn("w:fill"),
                ["E8EFF8", "EDF5F1", "FFF5E4", "F3EDFA", "FCEFF2"][depth % 5],
            )
            target._tc.get_or_add_tcPr().append(shade)
    if not scheme["classifications"]:
        paragraph(label("empty"))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
